import pytest
from pptx import Presentation as PptxPresentation

from app.schemas.slide_content import SLIDE_LAYOUTS, SlideContent
from app.tools.pptx_tool import build_pptx


@pytest.mark.parametrize("theme", ["classic", "modern", "warm"])
@pytest.mark.parametrize("include_sources", [False, True])
def test_build_pptx_preserves_layouts_titles_and_notes(
    tmp_path, sample_deck, theme, include_sources
):
    # Ensure the sample really covers every public layout.
    assert len(sample_deck.slides) == len(SLIDE_LAYOUTS)
    assert {slide.layout_type for slide in sample_deck.slides} == set(
        SLIDE_LAYOUTS
    )

    path = tmp_path / "presentation.pptx"

    build_pptx(
        sample_deck,
        path,
        theme=theme,
        include_sources_slide=include_sources,
    )

    assert path.is_file()
    reopened = PptxPresentation(path)

    # This sample contains sources, so enabling the option adds one slide.
    expected_count = len(sample_deck.slides) + int(include_sources)
    assert len(reopened.slides) == expected_count

    for expected, actual in zip(sample_deck.slides, reopened.slides):
        assert actual.shapes.title.text == expected.title
        assert (
            actual.notes_slide.notes_text_frame.text
            == expected.speaker_notes
        )

        if expected.layout_type == "chart":
            chart = next(
                shape.chart for shape in actual.shapes if shape.has_chart
            )
            assert [category.label for category in chart.plots[0].categories] == (
                expected.chart.categories
            )
            assert [series.name for series in chart.series] == [
                series.name for series in expected.chart.series
            ]
            for actual_series, expected_series in zip(
                chart.series, expected.chart.series
            ):
                assert list(actual_series.values) == expected_series.values

        if expected.layout_type == "table":
            table = next(
                shape.table for shape in actual.shapes if shape.has_table
            )
            assert len(table.rows) == len(expected.table.rows) + 1
            assert len(table.columns) == len(expected.table.columns)
            assert [
                table.cell(0, index).text
                for index in range(len(table.columns))
            ] == expected.table.columns

    if include_sources:
        assert reopened.slides[-1].shapes.title.text == "Sources"


def test_final_bullet_slide_keeps_title_and_notes(
    tmp_path, sample_deck
):
    # The renderer gives the last bullet slide its internal closing layout.
    sample_deck.slides.append(
        SlideContent(
            slide_number=6,
            title="Conclusion",
            layout_type="bullet_points",
            bullets=[{"text": "Focus on the strongest regions."}],
            speaker_notes="Summarize the recommended next steps.",
        )
    )

    path = tmp_path / "closing.pptx"
    build_pptx(sample_deck, path, include_sources_slide=False)

    reopened = PptxPresentation(path)

    assert len(reopened.slides) == 6
    assert reopened.slides[-1].shapes.title.text == "Conclusion"
    assert reopened.slides[-1].notes_slide.notes_text_frame.text == (
        "Summarize the recommended next steps."
    )