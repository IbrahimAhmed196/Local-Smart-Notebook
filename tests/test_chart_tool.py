import json

import pandas as pd
import pytest

from app.schemas.slide_content import ChartSpec, TableSpec
from app.tools.chart_tool import (
    MAX_BAR_CATEGORIES,
    MAX_LINE_POINTS,
    MAX_PIE_SLICES,
    MAX_TABLE_COLUMNS,
    MAX_TABLE_ROWS,
    ChartToolError,
    DataSource,
    build_chart_data,
    build_table_data,
    describe_sources,
)

SALES = pd.DataFrame(
    {
        "Region": ["North", "South", "North", "East", "South"],
        "Revenue": [100, 200, 300, 400, 500],
        "Units": [1, 2, 3, 4, 5],
        "Month": pd.to_datetime(["2024-01-01", "2024-02-01", "2024-03-01", "2024-04-01", "2024-05-01"]),
    }
)


@pytest.fixture()
def xlsx(tmp_path):
    path = tmp_path / "sales.xlsx"
    with pd.ExcelWriter(path) as writer:
        SALES.to_excel(writer, sheet_name="Q1", index=False)
        pd.DataFrame({"Item": ["a", "b"], "Cost": [7, 9]}).to_excel(writer, sheet_name="Costs", index=False)
    return [DataSource(document_id=1, filename="sales.xlsx", path=path)]


def chart(**kw):
    base = dict(chart_type="bar", title="T", document_id=1, sheet="Q1", x_column="Region", y_columns=["Revenue"])
    return ChartSpec(**{**base, **kw})


def test_describe_sources_lists_sheets_and_column_kinds_but_never_values(xlsx):
    (source,) = describe_sources(xlsx)
    assert source.document_id == 1 and source.filename == "sales.xlsx"
    assert [s.name for s in source.sheets] == ["Q1", "Costs"]
    assert source.sheets[0].columns == [
        ("Region", "text"), ("Revenue", "number"), ("Units", "number"), ("Month", "date"),
    ]
    assert "500" not in repr(source) and "North" not in repr(source)


def test_describe_sources_reads_semicolon_csv_and_skips_unreadable_and_other_files(tmp_path):
    csv = tmp_path / "a.csv"
    csv.write_text("city;sales\nCairo;10\nGiza;20\n", encoding="utf-8")
    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"not a workbook")
    txt = tmp_path / "notes.txt"
    txt.write_text("hello")
    sources = [DataSource(1, "a.csv", csv), DataSource(2, "bad.xlsx", bad), DataSource(3, "notes.txt", txt)]
    catalog = describe_sources(sources)
    assert [c.document_id for c in catalog] == [1]
    assert catalog[0].sheets[0].name is None
    assert catalog[0].sheets[0].columns == [("city", "text"), ("sales", "number")]


@pytest.mark.parametrize(
    "aggregation, expected",
    [
        ("sum", {"North": 400, "South": 700, "East": 400}),
        ("mean", {"North": 200, "South": 350, "East": 400}),
        ("count", {"North": 2, "South": 2, "East": 1}),
    ],
)
def test_aggregations_are_computed_from_the_file(xlsx, aggregation, expected):
    data = build_chart_data(chart(aggregation=aggregation), xlsx)
    assert dict(zip(data.categories, data.series[0].values)) == {k: float(v) for k, v in expected.items()}
    assert data.source == "sales.xlsx / Q1"


def test_aggregation_none_uses_the_rows_as_they_are(xlsx):
    data = build_chart_data(chart(aggregation="none", chart_type="line", x_column="Month"), xlsx)
    assert data.categories == ["2024-01-01", "2024-02-01", "2024-03-01", "2024-04-01", "2024-05-01"]
    assert data.series[0].values == [100.0, 200.0, 300.0, 400.0, 500.0]


def test_several_series_and_names_in_a_different_case(xlsx):
    data = build_chart_data(chart(sheet="q1", x_column="region", y_columns=["REVENUE", " units "]), xlsx)
    assert [s.name for s in data.series] == ["Revenue", "Units"]
    assert data.series[1].values == [4.0, 7.0, 4.0]


def test_the_default_sheet_is_the_first_one_and_a_csv_needs_no_sheet(xlsx, tmp_path):
    assert build_chart_data(chart(sheet=None), xlsx).source == "sales.xlsx / Q1"
    csv = tmp_path / "c.csv"
    csv.write_text("k,v\na,1\nb,2\n")
    data = build_chart_data(chart(sheet="ignored", x_column="k", y_columns=["v"]), [DataSource(1, "c.csv", csv)])
    assert data.categories == ["a", "b"] and data.source == "c.csv"


@pytest.mark.parametrize(
    "override",
    [
        dict(document_id=99),
        dict(sheet="Nope"),
        dict(x_column="Nope"),
        dict(y_columns=["Nope"]),
        dict(y_columns=["Region"]),  # text column: no numbers
        dict(y_columns=["Region"], x_column="Region"),  # nothing left once x is removed
    ],
)
def test_a_spec_that_does_not_match_the_file_is_rejected_not_guessed(xlsx, override):
    with pytest.raises(ChartToolError):
        build_chart_data(chart(**override), xlsx)


def test_count_works_on_a_text_column(xlsx):
    data = build_chart_data(chart(aggregation="count", y_columns=["Region"], x_column="Month"), xlsx)
    assert data.series[0].values == [1.0] * 5


def test_big_charts_keep_the_biggest_categories_and_pies_use_one_series(tmp_path):
    csv = tmp_path / "big.csv"
    pd.DataFrame({"k": [f"c{i}" for i in range(30)], "v": list(range(30)), "w": [1] * 30}).to_csv(csv, index=False)
    src = [DataSource(1, "big.csv", csv)]
    bar = build_chart_data(chart(sheet=None, x_column="k", y_columns=["v"]), src)
    assert len(bar.categories) == MAX_BAR_CATEGORIES and bar.categories[0] == "c29"
    pie = build_chart_data(chart(chart_type="pie", sheet=None, x_column="k", y_columns=["v", "w"]), src)
    assert len(pie.categories) == MAX_PIE_SLICES and len(pie.series) == 1
    line = build_chart_data(chart(chart_type="line", aggregation="none", sheet=None, x_column="k", y_columns=["v"]), src)
    assert len(line.categories) == MAX_LINE_POINTS and line.categories[0] == "c0"


def test_missing_numbers_are_dropped_never_replaced_by_zero(tmp_path):
    csv = tmp_path / "gaps.csv"
    csv.write_text("k,v\na,1\nb,\nc,3\n")
    data = build_chart_data(chart(sheet=None, x_column="k", y_columns=["v"], aggregation="none"), [DataSource(1, "gaps.csv", csv)])
    assert data.categories == ["a", "c"] and data.series[0].values == [1.0, 3.0]


def test_a_pie_cannot_show_negative_values(tmp_path):
    csv = tmp_path / "neg.csv"
    csv.write_text("k,v\na,5\nb,-2\n")
    with pytest.raises(ChartToolError):
        build_chart_data(chart(chart_type="pie", sheet=None, x_column="k", y_columns=["v"]), [DataSource(1, "neg.csv", csv)])


def test_chart_and_table_data_are_plain_json(xlsx):
    chart_json = build_chart_data(chart(aggregation="none", x_column="Month"), xlsx).model_dump()
    table_json = build_table_data(TableSpec(title="t", document_id=1, sheet="Q1", columns=["Region", "Month", "Revenue"]), xlsx).model_dump()
    json.dumps(chart_json)
    json.dumps(table_json)  # fails if a numpy / pandas type slipped through


def test_table_reads_the_real_cells(xlsx):
    spec = TableSpec(title="Rows", document_id=1, sheet="Q1", columns=["region", "Revenue", "Month"], max_rows=2)
    data = build_table_data(spec, xlsx)
    assert data.columns == ["Region", "Revenue", "Month"]
    assert data.rows == [["North", 100, "2024-01-01"], ["South", 200, "2024-02-01"]]
    assert data.source == "sales.xlsx / Q1"


def test_table_keeps_empty_cells_as_none_and_whole_numbers_as_numbers(tmp_path):
    csv = tmp_path / "t.csv"
    csv.write_text("name,score,ratio\nann,90,0.5\nbob,,2.25\n")
    data = build_table_data(TableSpec(title="t", document_id=1, columns=["name", "score", "ratio"]), [DataSource(1, "t.csv", csv)])
    assert data.rows == [["ann", 90, 0.5], ["bob", None, 2.25]]


def test_table_limits_rows_and_columns(tmp_path):
    csv = tmp_path / "wide.csv"
    pd.DataFrame({f"c{i}": range(40) for i in range(10)}).to_csv(csv, index=False)
    spec = TableSpec(title="t", document_id=1, columns=[f"c{i}" for i in range(10)], max_rows=500)
    data = build_table_data(spec, [DataSource(1, "wide.csv", csv)])
    assert len(data.columns) == MAX_TABLE_COLUMNS and len(data.rows) == MAX_TABLE_ROWS
    assert len(build_table_data(spec.model_copy(update={"max_rows": 0}), [DataSource(1, "wide.csv", csv)]).rows) == 1


def test_table_with_an_unknown_column_is_rejected(xlsx):
    with pytest.raises(ChartToolError):
        build_table_data(TableSpec(title="t", document_id=1, sheet="Q1", columns=["Region", "Nope"]), xlsx)


# ---- Week 6 presentation edits (DAT-04): robust spreadsheet reading and caching ----

import openpyxl

from app.tools import chart_tool
from app.tools.chart_tool import _parse_number, clear_sheet_cache


@pytest.fixture(autouse=True)
def _fresh_cache():
    clear_sheet_cache()
    yield
    clear_sheet_cache()


def _workbook(tmp_path, rows, name="report.xlsx", merge_title=False):
    path = tmp_path / name
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Report"
    for row in rows:
        sheet.append(row)
    if merge_title:
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)
    book.save(path)
    return [DataSource(1, name, path)]


def test_a_title_row_above_the_header_is_skipped(tmp_path):
    src = _workbook(
        tmp_path,
        [["Q1 sales report", None, None], [None, None, None], ["Region", "Revenue", "Units"], ["North", 100, 1], ["South", 200, 2]],
        merge_title=True,
    )
    (cat,) = describe_sources(src)
    assert cat.sheets[0].columns == [("Region", "text"), ("Revenue", "number"), ("Units", "number")]
    data = build_chart_data(chart(sheet="Report"), src)
    assert data.categories == ["North", "South"] and data.series[0].values == [100.0, 200.0]


def test_a_normal_header_in_row_one_still_works_when_a_data_row_is_full_and_the_header_is_not(tmp_path):
    src = _workbook(tmp_path, [["Region", "Revenue", None], ["North", 100, 5]])
    data = build_table_data(TableSpec(title="t", document_id=1, columns=["Region", "Revenue"]), src)
    assert data.rows == [["North", 100]]


@pytest.mark.parametrize(
    "text, expected",
    [
        ("$1,200", 1200.0),
        ("€ 1 200", 1200.0),
        ("-$1,200.50", -1200.5),
        ("(300)", -300.0),
        ("12%", 0.12),
        ("1.234,5", 1234.5),
        ("1,234.5", 1234.5),
        ("12,5", 12.5),
        ("1.234.567", 1234567.0),
        (7, 7.0),
        ("  42 ", 42.0),
    ],
)
def test_numbers_written_as_text_are_understood(text, expected):
    assert _parse_number(text) == pytest.approx(expected)


@pytest.mark.parametrize("text", ["North", "Q1 2024", "2024-01-05", "abc1", "", "1,2,3x", None, float("nan"), True])
def test_text_that_is_not_a_number_is_not_one(text):
    assert _parse_number(text) is None


def test_currency_and_percent_text_columns_can_be_charted_and_are_listed_as_numbers(tmp_path):
    src = _workbook(
        tmp_path,
        [["Region", "Revenue", "Growth", "Day"],
         ["North", "$1,200", "12%", "2024-01-05"],
         ["South", "$800", "5%", "2024-01-06"],
         ["North", "$300", "7%", "2024-01-07"]],
    )
    (cat,) = describe_sources(src)
    assert cat.sheets[0].columns == [("Region", "text"), ("Revenue", "number"), ("Growth", "number"), ("Day", "date")]
    data = build_chart_data(chart(sheet="Report", y_columns=["Revenue", "Growth"]), src)
    assert dict(zip(data.categories, data.series[0].values)) == {"North": 1500.0, "South": 800.0}
    assert data.series[1].values == pytest.approx([0.19, 0.05])
    table = build_table_data(TableSpec(title="t", document_id=1, sheet="Report", columns=["Revenue"], max_rows=1), src)
    assert table.rows == [["$1,200"]]  # a table shows the cell as it is written in the file


def test_a_column_of_words_is_still_not_a_number(tmp_path):
    src = _workbook(tmp_path, [["Region", "Note"], ["North", "good"], ["South", "bad"]])
    with pytest.raises(ChartToolError, match="has no numbers"):
        build_chart_data(chart(sheet="Report", y_columns=["Note"]), src)


def test_formulas_that_were_never_calculated_give_a_clear_error(tmp_path, caplog):
    src = _workbook(tmp_path, [["Region", "Price", "Total"], ["North", 10, "=B2*2"], ["South", 20, "=B3*2"]])
    with pytest.raises(ChartToolError, match="formulas that were never calculated"):
        build_chart_data(chart(sheet="Report", y_columns=["Total"]), src)
    with pytest.raises(ChartToolError, match="Column 'Total' is empty.*never calculated"):
        build_table_data(TableSpec(title="t", document_id=1, sheet="Report", columns=["Region", "Total"]), src)
    assert "formulas without a saved result" in caplog.text
    # the columns that hold real values still work
    assert build_chart_data(chart(sheet="Report", y_columns=["Price"]), src).series[0].values == [10.0, 20.0]


def test_a_file_is_read_once_per_run_and_again_after_it_changes(tmp_path, monkeypatch):
    src = _workbook(tmp_path, [["Region", "Revenue"], ["North", 1]])
    reads = []
    real = chart_tool.pd.read_excel
    monkeypatch.setattr(chart_tool.pd, "read_excel", lambda *a, **k: (reads.append(1), real(*a, **k))[1])

    describe_sources(src)
    build_chart_data(chart(sheet="Report"), src)
    build_chart_data(chart(sheet="Report", aggregation="mean"), src)
    build_table_data(TableSpec(title="t", document_id=1, sheet="Report", columns=["Region"]), src)
    assert len(reads) == 1

    book = openpyxl.load_workbook(src[0].path)
    book["Report"].append(["South", 2, ])
    book.save(src[0].path)  # a different file size -> the cache entry is not reused
    assert build_chart_data(chart(sheet="Report"), src).categories == ["North", "South"]
    assert len(reads) == 2

    clear_sheet_cache()
    describe_sources(src)
    assert len(reads) == 3


def test_the_agent_clears_the_cache_after_a_deck(tmp_path, scripted_llm):
    from app.agents.slide_content_agent import generate_presentation_content
    from tests.slide_samples import mock_outline

    src = _workbook(tmp_path, [["Region", "Revenue"], ["North", 1]])
    describe_sources(src)
    assert chart_tool._SHEET_CACHE
    # every scripted answer is a failure -> placeholder slides; the cache must be cleared anyway
    outline = mock_outline()
    client = scripted_llm(*[RuntimeError("x")] * len(outline.slides))
    generate_presentation_content(outline, data_files=src, client=client)
    assert not chart_tool._SHEET_CACHE
