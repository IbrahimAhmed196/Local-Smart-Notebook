"""Week 6 presentation edits (DAT-02): chart and table checks return a list of issues with a severity and a hint."""
import openpyxl
import pandas as pd
import pytest

from app.agents.slide_content_agent import generate_slide_content
from app.schemas.slide_content import Bullet, ChartSpec, SlideContentDraft, TableSpec
from app.schemas.slides import SlideOutlineItem
from app.tools import chart_tool
from app.tools.chart_tool import (
    MAX_BAR_CATEGORIES,
    MAX_CHART_SERIES,
    MAX_LINE_POINTS,
    MAX_PIE_SLICES,
    MAX_TABLE_COLUMNS,
    ChartToolError,
    CheckResult,
    DataSource,
    Issue,
    build_chart_data,
    check_chart,
    check_table,
    clear_sheet_cache,
)
from tests.slide_samples import FakeClient


@pytest.fixture(autouse=True)
def _fresh_cache():
    clear_sheet_cache()
    yield
    clear_sheet_cache()


def csv_source(tmp_path, frame, name="d.csv"):
    path = tmp_path / name
    frame.to_csv(path, index=False)
    return [DataSource(1, name, path)]


def chart(**kw):
    base = dict(chart_type="bar", title="T", document_id=1, x_column="k", y_columns=["v"], aggregation="sum")
    return ChartSpec(**{**base, **kw})


def table(**kw):
    base = dict(title="T", document_id=1, columns=["k", "v"], max_rows=5)
    return TableSpec(**{**base, **kw})


def issue(result: CheckResult, code: str) -> Issue:
    found = [i for i in result.issues if i.code == code]
    assert found, f"no '{code}' issue in {[i.code for i in result.issues]}"
    return found[0]


def small(tmp_path, n=4):
    return csv_source(tmp_path, pd.DataFrame({"k": [f"c{i}" for i in range(n)], "v": [10 * (i + 1) for i in range(n)]}))


# ---------------------------------------------------------------- the result itself

def test_a_clean_chart_and_table_have_no_issues(tmp_path):
    src = small(tmp_path)
    for result in (check_chart(chart(), src), check_table(table(), src)):
        assert result.ok and result.issues == [] and result.data is not None


def test_every_issue_has_a_severity_a_message_and_a_hint(tmp_path):
    src = csv_source(tmp_path, pd.DataFrame({"k": list("abc"), "v": [1, 2, 3], "w": [4, 5, 6]}))
    result = check_chart(chart(chart_type="pie", y_columns=["v", "w"]), src)  # two series on a pie
    first = result.issues[0]
    assert first.severity == "warning" and first.message and first.hint
    assert result.ok and result.warnings == [first] and result.errors == []  # a warning never stops the chart


def test_an_error_gives_no_data_and_is_listed_as_an_error(tmp_path):
    result = check_chart(chart(x_column="nope"), small(tmp_path))
    assert not result.ok and result.data is None
    assert [i.severity for i in result.issues] == ["error"]  # what GEN-06 can feed back to the model


def test_build_chart_data_still_raises_but_now_with_a_code_and_a_hint(tmp_path):
    with pytest.raises(ChartToolError, match="Column 'nope' not found") as caught:
        build_chart_data(chart(x_column="nope"), small(tmp_path))
    assert caught.value.code == "unknown_column" and "DATA FILES" in caught.value.hint
    assert caught.value.issue == Issue("unknown_column", "error", str(caught.value), caught.value.hint)


# ---------------------------------------------------------------- the file does not match the spec (errors)

def test_an_unknown_document_is_an_issue_not_an_exception(tmp_path):
    result = check_chart(chart(document_id=99), small(tmp_path))
    assert issue(result, "unknown_document").severity == "error"


def test_an_unknown_sheet_is_an_issue(tmp_path):
    path = tmp_path / "book.xlsx"
    pd.DataFrame({"k": ["a"], "v": [1]}).to_excel(path, sheet_name="Q1", index=False)
    result = check_chart(chart(sheet="Nope"), [DataSource(1, "book.xlsx", path)])
    assert "Q1" in issue(result, "unknown_sheet").message  # it lists the sheets that do exist


def test_an_unreadable_file_is_an_issue(tmp_path):
    path = tmp_path / "bad.xlsx"
    path.write_bytes(b"not a workbook")
    assert issue(check_chart(chart(), [DataSource(1, "bad.xlsx", path)]), "unreadable_file").severity == "error"
    assert issue(check_table(table(), [DataSource(1, "bad.xlsx", path)]), "unreadable_file").severity == "error"


def test_an_unknown_column_is_an_issue_for_charts_and_tables(tmp_path):
    src = small(tmp_path)
    assert issue(check_chart(chart(y_columns=["nope"]), src), "unknown_column").severity == "error"
    assert issue(check_table(table(columns=["k", "nope"]), src), "unknown_column").severity == "error"


def test_a_chart_needs_a_value_column_besides_x(tmp_path):
    result = check_chart(chart(y_columns=["k"]), small(tmp_path))  # the only y column is the x column
    assert issue(result, "no_value_column").severity == "error"


def test_a_column_without_numbers_is_an_error(tmp_path):
    src = csv_source(tmp_path, pd.DataFrame({"k": ["a", "b"], "v": ["good", "bad"]}))
    assert issue(check_chart(chart(), src), "no_numbers").severity == "error"


def test_nothing_left_after_removing_empty_cells_is_an_error(tmp_path):
    # each row misses one of its two values, so no row has both
    src = csv_source(tmp_path, pd.DataFrame({"k": list("abcd"), "v": [1, 2, None, None], "w": [None, None, 3, 4]}))
    result = check_chart(chart(y_columns=["v", "w"], aggregation="none"), src)
    assert issue(result, "empty_result").severity == "error"


# ---------------------------------------------------------------- bar charts

def test_a_bar_chart_with_too_many_categories_is_cut_with_a_warning(tmp_path):
    src = csv_source(tmp_path, pd.DataFrame({"k": [f"c{i}" for i in range(30)], "v": range(30)}))
    result = check_chart(chart(), src)
    cut = issue(result, "categories_truncated")
    assert cut.severity == "warning" and str(MAX_BAR_CATEGORIES) in cut.message
    assert len(result.data.categories) == MAX_BAR_CATEGORIES and result.data.categories[0] == "c29"  # the biggest stay


def test_an_x_column_that_looks_like_an_id_is_flagged(tmp_path):
    ids = pd.DataFrame({"k": [f"order-{i}" for i in range(30)], "v": range(30)})
    assert issue(check_chart(chart(), csv_source(tmp_path, ids)), "x_looks_like_id").severity == "warning"


def test_repeated_categories_dates_and_lines_are_not_flagged_as_ids(tmp_path):
    repeated = pd.DataFrame({"k": ["north", "south"] * 15, "v": range(30)})
    assert "x_looks_like_id" not in [i.code for i in check_chart(chart(), csv_source(tmp_path, repeated)).issues]
    dates = pd.DataFrame({"k": pd.date_range("2024-01-01", periods=30).astype(str), "v": range(30)})
    assert "x_looks_like_id" not in [i.code for i in check_chart(chart(), csv_source(tmp_path, dates, "dates.csv")).issues]
    ids = pd.DataFrame({"k": [f"order-{i}" for i in range(30)], "v": range(30)})
    line = check_chart(chart(chart_type="line", aggregation="none"), csv_source(tmp_path, ids, "ids.csv"))
    assert "x_looks_like_id" not in [i.code for i in line.issues]


def test_aggregation_none_with_repeated_x_values_is_flagged_but_sum_is_not(tmp_path):
    src = csv_source(tmp_path, pd.DataFrame({"k": ["a", "a", "b"], "v": [1, 2, 3]}))
    duplicate = issue(check_chart(chart(aggregation="none"), src), "duplicate_x")
    assert duplicate.severity == "warning" and "1 x value" in duplicate.message and "sum" in duplicate.hint
    assert "duplicate_x" not in [i.code for i in check_chart(chart(aggregation="sum"), src).issues]


def test_a_value_column_that_is_mostly_empty_is_flagged(tmp_path):
    src = csv_source(tmp_path, pd.DataFrame({"k": list("abcd"), "v": [5, None, None, None]}))
    result = check_chart(chart(aggregation="none"), src)
    mostly = issue(result, "column_mostly_empty")
    assert mostly.severity == "warning" and "1 of 4 rows" in mostly.message
    assert result.data.categories == ["a"]  # the empty rows were left out, never turned into 0


def test_a_value_column_with_text_in_some_rows_is_flagged(tmp_path):
    src = csv_source(tmp_path, pd.DataFrame({"k": list("abcd"), "v": ["10", "n/a", "tbd", "-"]}))
    assert issue(check_chart(chart(aggregation="none"), src), "column_mostly_empty").severity == "warning"


# ---------------------------------------------------------------- line charts

def test_a_line_chart_sorts_a_date_x_column(tmp_path):
    frame = pd.DataFrame({"k": ["2024-03-01", "2024-01-01", "2024-02-01"], "v": [3, 1, 2]})
    result = check_chart(chart(chart_type="line", aggregation="none"), csv_source(tmp_path, frame))
    assert result.data.categories == ["2024-01-01", "2024-02-01", "2024-03-01"]
    assert result.data.series[0].values == [1.0, 2.0, 3.0]  # the values moved with their dates
    assert result.issues == []


def test_a_line_chart_sorts_a_number_x_column(tmp_path):
    frame = pd.DataFrame({"year": [2022, 2020, 2021], "v": [30, 10, 20]})
    result = check_chart(chart(chart_type="line", aggregation="none", x_column="year"), csv_source(tmp_path, frame))
    assert result.data.categories == ["2020", "2021", "2022"] and result.data.series[0].values == [10.0, 20.0, 30.0]


def test_a_line_chart_on_text_categories_is_flagged_and_keeps_the_file_order(tmp_path):
    frame = pd.DataFrame({"k": ["south", "north", "east"], "v": [1, 2, 3]})
    result = check_chart(chart(chart_type="line", aggregation="none"), csv_source(tmp_path, frame))
    assert issue(result, "line_x_not_ordered").severity == "warning"
    assert result.data.categories == ["south", "north", "east"]


def test_a_long_line_is_cut_with_a_warning(tmp_path):
    frame = pd.DataFrame({"day": range(40), "v": range(40)})
    result = check_chart(chart(chart_type="line", aggregation="none", x_column="day"), csv_source(tmp_path, frame))
    assert issue(result, "points_truncated").severity == "warning"
    assert len(result.data.categories) == MAX_LINE_POINTS


def test_a_line_needs_at_least_two_points(tmp_path):
    src = csv_source(tmp_path, pd.DataFrame({"day": [1], "v": [5]}))
    assert issue(check_chart(chart(chart_type="line", aggregation="none", x_column="day"), src), "too_few_points").severity == "error"


# ---------------------------------------------------------------- pie charts

def test_a_pie_with_extra_series_uses_the_first_one_and_says_so(tmp_path):
    src = csv_source(tmp_path, pd.DataFrame({"k": list("abc"), "v": [1, 2, 3], "w": [4, 5, 6]}))
    result = check_chart(chart(chart_type="pie", y_columns=["v", "w"]), src)
    assert issue(result, "pie_extra_series").severity == "warning"
    assert [s.name for s in result.data.series] == ["v"]


def test_a_pie_with_too_many_slices_keeps_the_biggest_with_a_warning(tmp_path):
    src = csv_source(tmp_path, pd.DataFrame({"k": [f"c{i}" for i in range(15)], "v": range(1, 16)}))
    result = check_chart(chart(chart_type="pie"), src)
    assert issue(result, "categories_truncated").severity == "warning"
    assert len(result.data.categories) == MAX_PIE_SLICES


def test_a_pie_cannot_have_negative_values(tmp_path):
    src = csv_source(tmp_path, pd.DataFrame({"k": ["a", "b"], "v": [5, -2]}))
    assert issue(check_chart(chart(chart_type="pie"), src), "pie_negative").severity == "error"


def test_a_negative_value_is_found_even_when_the_pie_is_cut_to_its_biggest_slices(tmp_path):
    # 10 categories, the negative one is the smallest: cutting to 8 slices must not hide it
    values = [50, 40, 30, 20, 10, 9, 8, 7, 6, -5]
    src = csv_source(tmp_path, pd.DataFrame({"k": list("abcdefghij"), "v": values}))
    result = check_chart(chart(chart_type="pie"), src)
    assert not result.ok and issue(result, "pie_negative").severity == "error"
    assert result.data is None


def test_a_pie_total_must_be_above_zero(tmp_path):
    src = csv_source(tmp_path, pd.DataFrame({"k": ["a", "b"], "v": [0, 0]}))
    assert issue(check_chart(chart(chart_type="pie"), src), "pie_total_zero").severity == "error"


def test_a_pie_needs_at_least_two_slices(tmp_path):
    src = csv_source(tmp_path, pd.DataFrame({"k": ["a"], "v": [5]}))
    assert issue(check_chart(chart(chart_type="pie"), src), "too_few_points").severity == "error"


# ---------------------------------------------------------------- several series

def test_more_series_than_the_limit_are_cut_with_a_warning(tmp_path):
    columns = {f"s{i}": [1, 2, 3] for i in range(MAX_CHART_SERIES + 2)}
    src = csv_source(tmp_path, pd.DataFrame({"k": list("abc"), **columns}))
    result = check_chart(chart(y_columns=list(columns)), src)
    assert issue(result, "too_many_series").severity == "warning"
    assert len(result.data.series) == MAX_CHART_SERIES


def test_series_on_very_different_scales_are_flagged(tmp_path):
    src = csv_source(tmp_path, pd.DataFrame({"k": list("abc"), "big": [1_000_000, 2_000_000, 3_000_000], "small": [1, 2, 3]}))
    warning = issue(check_chart(chart(y_columns=["big", "small"]), src), "scale_mismatch")
    assert warning.severity == "warning" and "'small' will look flat" in warning.message


def test_series_of_a_similar_size_are_not_flagged(tmp_path):
    src = csv_source(tmp_path, pd.DataFrame({"k": list("abc"), "v": [10, 20, 30], "w": [5, 15, 25]}))
    assert check_chart(chart(y_columns=["v", "w"]), src).issues == []


def test_chart_data_that_fails_the_schema_checks_is_an_issue_not_a_crash(tmp_path):
    src = csv_source(tmp_path, pd.DataFrame({"k": list("abc"), "Sales": [1, 2, 3], "sales": [4, 5, 6]}))
    result = check_chart(chart(y_columns=["Sales", "sales"]), src)  # two series whose names differ only in case
    assert issue(result, "invalid_chart_data").severity == "error" and result.data is None


# ---------------------------------------------------------------- tables

def test_a_table_needs_at_least_one_column(tmp_path):
    assert issue(check_table(table(columns=[]), small(tmp_path)), "no_columns").severity == "error"


def test_a_table_with_too_many_columns_is_cut_with_a_warning(tmp_path):
    src = csv_source(tmp_path, pd.DataFrame({f"c{i}": range(5) for i in range(10)}))
    result = check_table(table(columns=[f"c{i}" for i in range(10)]), src)
    assert issue(result, "columns_truncated").severity == "warning"
    assert len(result.data.columns) == MAX_TABLE_COLUMNS


def test_asking_for_more_rows_than_fit_is_flagged(tmp_path, monkeypatch):
    monkeypatch.setattr(chart_tool, "MAX_TABLE_ROWS", 3)
    src = csv_source(tmp_path, pd.DataFrame({"k": list("abcdefgh"), "v": range(8)}))
    result = check_table(table(max_rows=8), src)
    assert issue(result, "rows_truncated").severity == "warning"
    assert len(result.data.rows) == 3


def test_a_table_column_that_is_empty_in_the_rows_shown_is_flagged(tmp_path):
    frame = pd.DataFrame({"k": list("abcdefgh"), "note": [None] * 6 + ["x", "y"]})
    result = check_table(table(columns=["k", "note"], max_rows=4), csv_source(tmp_path, frame))
    assert issue(result, "empty_column").severity == "warning" and "'note'" in result.issues[0].message
    assert result.ok


def test_a_table_with_no_rows_is_an_error(tmp_path):
    path = tmp_path / "empty.csv"
    path.write_text("k,v\n", encoding="utf-8")
    result = check_table(table(), [DataSource(1, "empty.csv", path)])
    assert not result.ok and result.data is None  # a header without rows: there is nothing to show


def test_an_excel_sheet_with_only_a_header_row_has_no_data(tmp_path):
    path = tmp_path / "header_only.xlsx"
    book = openpyxl.Workbook()
    book.active.append(["k", "v"])
    book.save(path)
    result = check_table(table(), [DataSource(1, "header_only.xlsx", path)])
    assert not result.ok and result.data is None
    assert issue(result, "no_data").severity == "error" and issue(result, "no_data").hint


def test_a_column_of_formulas_that_were_never_calculated_is_an_error(tmp_path):
    path = tmp_path / "f.xlsx"
    book = openpyxl.Workbook()
    book.active.title = "Report"
    for row in [["k", "price", "total"], ["a", 10, "=B2*2"], ["b", 20, "=B3*2"]]:
        book.active.append(row)
    book.save(path)
    result = check_table(table(columns=["k", "total"], sheet="Report"), [DataSource(1, "f.xlsx", path)])
    assert issue(result, "formulas_not_calculated").severity == "error"
    assert "never calculated" in result.errors[0].message


def test_very_long_cells_are_flagged(tmp_path):
    frame = pd.DataFrame({"k": ["x" * 120, "short"], "v": [1, 2]})
    result = check_table(table(), csv_source(tmp_path, frame))
    long = issue(result, "long_cells")
    assert long.severity == "warning" and "1 cell" in long.message
    assert result.data.rows[0][0] == "x" * 120  # the cell itself is never changed


def test_table_data_that_fails_the_schema_checks_is_an_issue_not_a_crash(tmp_path):
    src = csv_source(tmp_path, pd.DataFrame({"k": ["a", "b"], "v": [1.0, float("inf")]}))
    result = check_table(table(), src)
    assert issue(result, "invalid_table_data").severity == "error" and result.data is None


# ---------------------------------------------------------------- what the slide writer does with the issues

def item(layout="chart"):
    return SlideOutlineItem(slide_number=4, title="Data", purpose="Show the data", layout_type=layout)


def answer(text, **kw):
    return SlideContentDraft(bullets=[Bullet(text=text, source_ids=[])], **kw)


def write(client, src):
    return generate_slide_content(item(), presentation_title="Deck", audience="Team", data_files=src, client=client)


def test_chart_warnings_become_slide_warnings_and_the_chart_is_kept(tmp_path):
    src = csv_source(tmp_path, pd.DataFrame({"k": [f"order-{i}" for i in range(30)], "v": range(30)}))
    first = answer("Orders", chart=chart())
    slide = write(FakeClient.sequence(first, answer("Biggest order")), src)
    assert slide.layout_type == "chart" and slide.chart is not None
    assert any("looks like an id" in w and "Choose a column with fewer" in w for w in slide.warnings)  # message + hint
    assert any("only the 12 biggest" in w for w in slide.warnings)


def test_a_chart_error_drops_the_chart_and_the_warning_carries_the_hint(tmp_path):
    src = small(tmp_path)
    slide = write(FakeClient.sequence(answer("Revenue", chart=chart(y_columns=["profit"]))), src)
    assert slide.chart is None and slide.layout_type == "bullet_points"
    assert any(w.startswith("Slide 4") and "Chart dropped: Column 'profit' not found" in w and "Copy the column name" in w for w in slide.warnings)


def test_table_warnings_become_slide_warnings_and_a_table_error_drops_the_table(tmp_path):
    src = csv_source(tmp_path, pd.DataFrame({"k": ["x" * 120, "short"], "v": [1, 2]}))
    slide = generate_slide_content(
        item("table"), presentation_title="Deck", audience="Team", data_files=src,
        client=FakeClient.sequence(answer("Rows", table=table()), answer("Two rows")),
    )
    assert slide.table is not None and any("longer than" in w for w in slide.warnings)

    broken = generate_slide_content(
        item("table"), presentation_title="Deck", audience="Team", data_files=src,
        client=FakeClient.sequence(answer("Rows", table=table(columns=["k", "nope"]))),
    )
    assert broken.table is None and any("Table dropped: Column 'nope' not found" in w for w in broken.warnings)


# ---------------------------------------------------------------- the new limits are settings (.env), like the others (ARC-07)

def test_the_new_settings_have_the_defaults_the_docs_promise(monkeypatch):
    from app.core.config import Settings

    s = Settings(_env_file=None)
    assert (s.chart_max_series, s.chart_scale_ratio_warning, s.table_max_cell_chars) == (4, 100, 60)
    assert (s.catalog_example_labels, s.slide_data_facts) == (5, True)


def test_the_new_settings_can_be_set_in_the_environment(monkeypatch):
    from app.core.config import Settings

    for name, value in dict(
        CHART_MAX_SERIES="6", CHART_SCALE_RATIO_WARNING="50", TABLE_MAX_CELL_CHARS="80",
        CATALOG_EXAMPLE_LABELS="0", SLIDE_DATA_FACTS="false",
    ).items():
        monkeypatch.setenv(name, value)
    s = Settings(_env_file=None)
    assert (s.chart_max_series, s.chart_scale_ratio_warning, s.table_max_cell_chars) == (6, 50, 80)
    assert (s.catalog_example_labels, s.slide_data_facts) == (0, False)


@pytest.mark.parametrize(
    "name, value",
    [("CHART_MAX_SERIES", "0"), ("CHART_SCALE_RATIO_WARNING", "1"), ("TABLE_MAX_CELL_CHARS", "2"), ("CATALOG_EXAMPLE_LABELS", "-1")],
)
def test_a_nonsense_value_for_a_new_setting_stops_the_app_at_start(monkeypatch, name, value):
    from pydantic import ValidationError

    from app.core.config import Settings

    monkeypatch.setenv(name, value)
    with pytest.raises(ValidationError, match=name.lower()):
        Settings(_env_file=None)
