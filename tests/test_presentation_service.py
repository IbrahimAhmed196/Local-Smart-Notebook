import pytest
from pptx import Presentation as PptxPresentation

from app.core.exceptions import ConflictError
from app.database.models import Notebook, Presentation, PresentationStatus
from app.schemas.presentations import PresentationCreate
from app.schemas.slides import PresentationOutlineResponse, SlideOutlineItem
from app.services import presentation_service as service
from app.presentations.outline import clean_outline
from app.presentations import runner


@pytest.fixture()
def notebook(db, notebook_id):
    return db.get(Notebook, notebook_id)


def create_presentation(db, notebook):
    return service.create_presentation(
        db,
        notebook,
        PresentationCreate(
            topic="Sales review",
            num_slides=6,
            include_sources_slide=True,
        ),
    )


@pytest.mark.parametrize("fails", [False, True])
def test_generation_lifecycle(
    db, notebook, fake_generator, monkeypatch, fails
):
    events = []

    # Observe PENDING immediately before the real inline job starts.
    original_submit = service.job_manager.submit

    def observe_submit(fn, presentation_id):
        db.expire_all()
        presentation = db.get(Presentation, presentation_id)
        events.append((presentation.status, presentation.stage))
        return original_submit(fn, presentation_id)

    monkeypatch.setattr(service.job_manager, "submit", observe_submit)

    # Record committed transitions while preserving the real _touch().
    original_touch = runner._touch

    def observe_touch(session, presentation, **fields):
        original_touch(session, presentation, **fields)
        events.append((presentation.status, presentation.stage))

    monkeypatch.setattr(runner, "_touch", observe_touch)

    if fails:
        fake_generator.error = RuntimeError("Generation failed")

    presentation = create_presentation(db, notebook)

    assert events[0] == (PresentationStatus.PENDING, "queued")
    assert events[1] == (PresentationStatus.RUNNING, "starting")
    assert (PresentationStatus.RUNNING, "outline") in events

    assert len(fake_generator.calls) == 1
    assert fake_generator.calls[0]["notebook_id"] == notebook.id

    if fails:
        assert presentation.status == PresentationStatus.FAILED
        assert presentation.error == "RuntimeError: Generation failed"
        assert presentation.file_path is None
        assert presentation.deck is None
        assert events[-1][0] == PresentationStatus.FAILED
    else:
        assert events[-1] == (PresentationStatus.COMPLETED, "done")
        assert presentation.error is None
        assert presentation.deck is not None
        assert service.get_download(presentation).is_file()


@pytest.mark.parametrize(
    "initial_status",
    [PresentationStatus.COMPLETED, PresentationStatus.FAILED],
)
def test_update_deck_renders_without_generating_again(
    db, notebook, fake_generator, initial_status
):
    presentation = create_presentation(db, notebook)
    presentation.status = initial_status
    presentation.error = "Old failure"
    db.commit()

    edited = service.get_deck(presentation)
    edited.slides[0].title = "Edited sales review"
    edited.slides[0].speaker_notes = "Updated speaker notes."

    updated = service.update_deck(db, presentation, edited)

    assert updated.status == PresentationStatus.COMPLETED
    assert updated.stage == "done"
    assert updated.error is None
    assert updated.deck == edited.model_dump(mode="json")

    # Only the original creation called the generator.
    assert len(fake_generator.calls) == 1

    reopened = PptxPresentation(service.get_download(updated))
    assert reopened.slides[0].shapes.title.text == "Edited sales review"
    assert (
        reopened.slides[0].notes_slide.notes_text_frame.text
        == "Updated speaker notes."
    )


@pytest.mark.parametrize(
    "status",
    [PresentationStatus.PENDING, PresentationStatus.RUNNING],
)
def test_update_deck_rejects_active_jobs(
    db, notebook, sample_deck, status
):
    presentation = Presentation(
        notebook_id=notebook.id,
        title="Busy presentation",
        status=status,
        options={},
    )
    db.add(presentation)
    db.commit()

    with pytest.raises(ConflictError, match="still being generated"):
        service.update_deck(db, presentation, sample_deck)


def test_delete_removes_database_row_and_file(
    db, notebook, fake_generator
):
    presentation = create_presentation(db, notebook)
    presentation_id = presentation.id
    path = service.get_download(presentation)

    assert path.is_file()

    service.delete_presentation(db, presentation)

    assert db.get(Presentation, presentation_id) is None
    assert not path.exists()


def test_fail_interrupted_jobs_only_changes_unfinished_jobs(db, notebook):
    statuses = [
        PresentationStatus.PENDING,
        PresentationStatus.RUNNING,
        PresentationStatus.COMPLETED,
        PresentationStatus.FAILED,
    ]
    presentations = [
        Presentation(
            notebook_id=notebook.id,
            title=status,
            status=status,
            options={},
            error="Existing failure"
            if status == PresentationStatus.FAILED
            else None,
        )
        for status in statuses
    ]
    db.add_all(presentations)
    db.commit()

    assert runner.fail_interrupted_jobs() == 2

    # The function uses a separate database session.
    db.expire_all()

    for presentation in presentations[:2]:
        assert presentation.status == PresentationStatus.FAILED
        assert presentation.error == "Interrupted by a server restart."

    assert presentations[2].status == PresentationStatus.COMPLETED
    assert presentations[2].error is None
    assert presentations[3].status == PresentationStatus.FAILED
    assert presentations[3].error == "Existing failure"

    # Calling it again should not change already-finished jobs.
    assert runner.fail_interrupted_jobs() == 0


def make_outline():
    return PresentationOutlineResponse(
        presentation_title="Model title",
        target_audience="",
        slides=[
            SlideOutlineItem(
                slide_number=30,
                title="Table",
                purpose="Show records",
                layout_type="table",
            ),
            SlideOutlineItem(
                slide_number=10,
                title="Introduction",
                purpose="Introduce the topic",
                layout_type="bullet_points",
            ),
            SlideOutlineItem(
                slide_number=20,
                title="Chart",
                purpose="Compare revenue",
                layout_type=" Chart ",
            ),
            SlideOutlineItem(
                slide_number=40,
                title="Extra slide",
                purpose="This should be trimmed",
                layout_type="bullet_points",
            ),
        ],
    )


@pytest.mark.parametrize("has_data", [False, True])
@pytest.mark.parametrize("include_sources", [False, True])
def test_clean_outline_repairs_numbering_layouts_and_title(
    has_data, include_sources
):
    outline = make_outline()
    original = outline.model_dump()

    options = {
        "title": "User title",
        "audience": "Managers",
        # Both cases request three content slides.
        "num_slides": 4 if include_sources else 3,
        "include_sources_slide": include_sources,
    }

    cleaned = clean_outline(outline, options, has_data)

    assert cleaned.presentation_title == "User title"
    assert cleaned.target_audience == "Managers"
    assert [slide.slide_number for slide in cleaned.slides] == [1, 2, 3]
    assert [slide.title for slide in cleaned.slides] == [
        "User title",
        "Chart",
        "Table",
    ]
    assert [slide.layout_type for slide in cleaned.slides] == (
        ["title", "chart", "table"]
        if has_data
        else ["title", "bullet_points", "bullet_points"]
    )

    # Cleaning should return repaired objects without modifying the input.
    assert outline.model_dump() == original


@pytest.mark.parametrize(
    ("count", "message"),
    [
        (0, "without any slides"),
        (1, "Expected 3 outline slides, received 1"),
    ],
)
def test_clean_outline_rejects_empty_or_short_outlines(count, message):
    outline = make_outline()
    outline.slides = outline.slides[:count]

    options = {
        "num_slides": 4,
        "include_sources_slide": True,
        "audience": "Managers",
    }

    with pytest.raises(ValueError, match=message):
        clean_outline(outline, options, has_data=False)