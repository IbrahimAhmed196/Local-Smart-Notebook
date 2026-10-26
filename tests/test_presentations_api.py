from io import BytesIO

import pytest
from pptx import Presentation as PptxPresentation

from app.database.models import Presentation, PresentationStatus
from app.services import presentation_service as service


@pytest.fixture()
def created_presentation(client, auth, notebook_id, fake_generator):
    response = client.post(
        f"/notebooks/{notebook_id}/presentations",
        headers=auth,
        json={
            "topic": "Sales review",
            "num_slides": 6,
            "include_sources_slide": True,
        },
    )
    assert response.status_code == 202, response.text
    return response.json()


def test_create_poll_edit_download_delete(
    client, auth, notebook_id, created_presentation, fake_generator
):
    presentation_id = created_presentation["id"]
    url = f"/presentations/{presentation_id}"

    # Inline jobs finish before the create request returns.
    assert created_presentation["status"] == PresentationStatus.COMPLETED

    response = client.get(url, headers=auth)
    assert response.status_code == 200
    assert response.json()["stage"] == "done"
    assert response.json()["download_url"].endswith(f"{url}/download")

    response = client.get(
        f"/notebooks/{notebook_id}/presentations",
        headers=auth,
    )
    assert response.status_code == 200
    assert [item["id"] for item in response.json()] == [presentation_id]

    response = client.get(f"{url}/deck", headers=auth)
    assert response.status_code == 200
    deck = response.json()
    deck["slides"][0]["title"] = "Edited through the API"
    deck["slides"][0]["speaker_notes"] = "API edit notes."

    response = client.put(f"{url}/deck", headers=auth, json=deck)
    assert response.status_code == 200
    assert response.json()["status"] == PresentationStatus.COMPLETED

    response = client.get(f"{url}/deck", headers=auth)
    assert response.status_code == 200
    assert response.json() == deck
    assert len(fake_generator.calls) == 1

    response = client.get(f"{url}/download", headers=auth)
    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument."
        "presentationml.presentation"
    )
    assert "attachment" in response.headers["content-disposition"]

    reopened = PptxPresentation(BytesIO(response.content))
    assert reopened.slides[0].shapes.title.text == "Edited through the API"
    assert (
        reopened.slides[0].notes_slide.notes_text_frame.text
        == "API edit notes."
    )

    response = client.delete(url, headers=auth)
    assert response.status_code == 204
    assert response.content == b""

    assert client.get(url, headers=auth).status_code == 404
    assert client.get(f"{url}/deck", headers=auth).status_code == 404
    assert client.get(f"{url}/download", headers=auth).status_code == 404


def test_failed_generation_is_visible_when_polled(
    client, auth, notebook_id, fake_generator
):
    fake_generator.error = RuntimeError("Generation failed")

    response = client.post(
        f"/notebooks/{notebook_id}/presentations",
        headers=auth,
        json={"topic": "Sales review", "num_slides": 6},
    )
    assert response.status_code == 202
    url = f"/presentations/{response.json()['id']}"

    response = client.get(url, headers=auth)
    assert response.status_code == 200
    assert response.json()["status"] == PresentationStatus.FAILED
    assert response.json()["error"] == "RuntimeError: Generation failed"
    assert response.json()["download_url"] is None

    assert client.get(f"{url}/deck", headers=auth).status_code == 409
    assert client.get(f"{url}/download", headers=auth).status_code == 409


@pytest.mark.parametrize(
    "status",
    [PresentationStatus.PENDING, PresentationStatus.RUNNING],
)
def test_active_presentation_returns_conflicts(
    client, auth, notebook_id, db, sample_deck, status
):
    presentation = Presentation(
        notebook_id=notebook_id,
        title="Busy presentation",
        status=status,
        options={},
    )
    db.add(presentation)
    db.commit()
    url = f"/presentations/{presentation.id}"

    assert client.get(url, headers=auth).json()["status"] == status
    assert client.get(f"{url}/deck", headers=auth).status_code == 409
    assert client.get(f"{url}/download", headers=auth).status_code == 409

    response = client.put(
        f"{url}/deck",
        headers=auth,
        json=sample_deck.model_dump(mode="json"),
    )
    assert response.status_code == 409


def test_other_users_cannot_access_presentations(
    client, auth, login, notebook_id, created_presentation, sample_deck
):
    other_auth = login("other@example.com")
    presentation_id = created_presentation["id"]

    # Existing private IDs and missing IDs must both return 404.
    for target_id in (presentation_id, presentation_id + 9999):
        url = f"/presentations/{target_id}"

        for method, suffix in [
            ("GET", ""),
            ("GET", "/deck"),
            ("PUT", "/deck"),
            ("GET", "/download"),
            ("DELETE", ""),
        ]:
            kwargs = {"headers": other_auth}
            if method == "PUT":
                kwargs["json"] = sample_deck.model_dump(mode="json")

            response = client.request(method, url + suffix, **kwargs)
            assert response.status_code == 404, response.text

    notebook_url = f"/notebooks/{notebook_id}/presentations"

    assert client.get(notebook_url, headers=other_auth).status_code == 404
    response = client.post(
        notebook_url,
        headers=other_auth,
        json={"topic": "Unauthorized creation"},
    )
    assert response.status_code == 404

    # The unsuccessful DELETE must not remove the owner's presentation.
    response = client.get(
        f"/presentations/{presentation_id}",
        headers=auth,
    )
    assert response.status_code == 200


def test_missing_artifact_returns_404(
    client, auth, created_presentation, db
):
    presentation = db.get(Presentation, created_presentation["id"])
    service.get_download(presentation).unlink()

    response = client.get(
        f"/presentations/{presentation.id}/download",
        headers=auth,
    )
    assert response.status_code == 404