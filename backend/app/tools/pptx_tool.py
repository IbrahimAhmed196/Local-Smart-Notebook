"""PowerPoint creation tool (Stage 2)."""
"""Turns Task 2's finished deck (PresentationContent) into an editable .pptx file.
 
No AI and no database here: a deck goes in, a file comes out.
"""
 
import re
from dataclasses import dataclass
from pathlib import Path
 
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
 
from app.schemas.slide_content import Citation, PresentationContent, SlideContent
from app.tools import pptx_charts
from app.tools.pptx_themes import Theme, get_theme
 
# ---- page geometry (inches, 16:9) ----
SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.75
BODY_TOP, BODY_H = 1.8, 4.85
MAX_TABLE_ROWS_PER_SLIDE = 8
LABELS = {  # the few words the renderer adds itself
    "en": {
        "sources": "Sources",
        "footer": "Sources:",
        "page": "page",
        "pages": "pages",
        "cont": "(cont.)",
        "audience": "For:",
    },
    "ar": {
        "sources": "\u0627\u0644\u0645\u0635\u0627\u062f\u0631",
        "footer": "\u0627\u0644\u0645\u0635\u0627\u062f\u0631:",
        "page": "\u0635\u0641\u062d\u0629",
        "pages": "\u0635\u0641\u062d\u0627\u062a",
        "cont": "(\u062a\u0627\u0628\u0639)",
        "audience": "\u0627\u0644\u062c\u0645\u0647\u0648\u0631:",
    },
}
 
# the names Task 2 uses, plus a few names the renderer also accepts
LAYOUT_ALIASES = {
    "bullets": "bullet_points",
    "bullet_point": "bullet_points",
    "bullet": "bullet_points",
    "conclusion": "closing",
    "two_column": "two_columns",
    "two_cols": "two_columns",
    "comparison": "two_columns",
}
KNOWN_LAYOUTS = {"title", "bullet_points", "two_columns", "chart", "table", "closing"}
 
 
# ---------------------------------------------------------------- helpers
def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)
 
 
def _mx(x: float, w: float, rtl: bool) -> float:
    """Mirror a horizontal position for right-to-left decks."""
    return SLIDE_W - x - w if rtl else x
 
 
def _style_run(run, size, color, font, bold=False, italic=False):
    f = run.font
    f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, font
    f.color.rgb = _rgb(color)
    rpr = run._r.get_or_add_rPr()  # complex-script font, so Arabic text keeps a real font
    if rpr.find(qn("a:cs")) is None:
        cs = etree.SubElement(rpr, qn("a:cs"))
        cs.set("typeface", "Arial")
        latin = rpr.find(qn("a:latin"))
        if latin is not None:
            latin.addnext(cs)
 
 
def _para(
    tf,
    first,
    text,
    size,
    color,
    font,
    rtl,
    bold=False,
    italic=False,
    bullet=False,
    align=None,
    space_after=8,
):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    run = p.add_run()
    run.text = text
    _style_run(run, size, color, font, bold, italic)
    p.space_after = Pt(space_after)
    p.alignment = align or (PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT)
    ppr = p._p.get_or_add_pPr()
    if rtl:
        ppr.set("rtl", "1")
    if bullet:
        indent = Inches(0.35)
        ppr.set("marR" if rtl else "marL", str(int(indent)))
        ppr.set("indent", str(-int(indent)))
        bu = etree.SubElement(ppr, qn("a:buChar"))
        bu.set("char", "\u2022")
    return p
 
 
def _textbox(slide, x, y, w, h, rtl, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(_mx(x, w, rtl)), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE  # safety net; sizes are chosen below
    return box, tf
 
 
def _rect(slide, x, y, w, h, fill, rtl, shape=MSO_SHAPE.RECTANGLE):
    shp = slide.shapes.add_shape(
        shape, Inches(_mx(x, w, rtl)), Inches(y), Inches(w), Inches(h)
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(fill)
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp
 
 
def _bullet_size(texts: list[str], base: int = 24) -> int:
    """Pick a font size from the amount of text, so it fits without overflowing."""
    n, total = len(texts), sum(len(t) for t in texts)
    size = base
    if n > 4 or total > 360:
        size = base - 4
    if n > 5 or total > 520:
        size = base - 6
    if total > 700:
        size = base - 8
    return max(size, 14)
 
 
def _fmt_cell(value) -> str:
    """Table cells arrive as str / int / float / None from Task 2's file reader."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return f"{value:,}"
    if isinstance(value, float):
        return f"{int(value):,}" if value.is_integer() else f"{value:,.2f}"
    return str(value)
 
 
def _is_number(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)
 
 
def detect_rtl(deck: PresentationContent) -> bool:
    """True when most letters in the deck are Arabic or Hebrew."""
    text = " ".join(
        [deck.presentation_title]
        + [s.title for s in deck.slides]
        + [b.text for s in deck.slides for b in s.bullets + s.right_bullets]
    )
    rtl = len(re.findall(r"[\u0590-\u08ff]", text))
    latin = len(re.findall(r"[A-Za-z]", text))
    return rtl > latin
 
 
def detect_language(deck: PresentationContent) -> str:
    """'ar' when most letters are Arabic (the added labels are then Arabic), else 'en'."""
    text = " ".join([deck.presentation_title] + [s.title for s in deck.slides])
    arabic = len(re.findall(r"[\u0600-\u06ff]", text))
    latin = len(re.findall(r"[A-Za-z]", text))
    return "ar" if arabic > latin else "en"
 
 
def _footer_text(slide: SlideContent, lang: str) -> str:
    """'Sources: report.pdf (pages 2, 7) | sales.csv' for the slide footer."""
    parts = _sources_lines([slide], lang)
    for data in (slide.chart, slide.table):
        if data is not None and data.source:
            parts.append(data.source)
    unique = list(dict.fromkeys(parts))
    if not unique:
        return ""
    line = f"{LABELS[lang]['footer']} " + " | ".join(unique)
    return line if len(line) <= 170 else line[:167] + "..."
 
 
# ---------------------------------------------------------------- planning
@dataclass
class Planned:
    """One slide as it will be drawn: the content, the layout actually used, the footer."""
 
    slide: SlideContent
    kind: str  # title | bullet_points | two_columns | chart | table | closing | sources
    footer: str = ""
 
 
def _layout_of(slide: SlideContent) -> str:
    raw = (slide.layout_type or "").strip().lower().replace(" ", "_").replace("-", "_")
    layout = LAYOUT_ALIASES.get(raw, raw)
    if layout not in KNOWN_LAYOUTS:
        layout = "bullet_points"  # the contract: an unknown layout is a bullet slide
    if layout == "chart" and slide.chart is None:
        layout = "bullet_points"
    if layout == "table" and slide.table is None:
        layout = "bullet_points"
    return layout
 
 
def _sources_lines(slides: list[SlideContent], lang: str = "en") -> list[str]:
    """One line per file: 'report.pdf (pages 2, 3, 4)'."""
    pages: dict[str, set[int]] = {}
    for s in slides:
        for c in s.citations:
            pages.setdefault(c.filename, set())
            if c.page:
                pages[c.filename].add(c.page)
    lines = []
    for name, nums in pages.items():
        if not nums:
            lines.append(name)
        else:
            label = LABELS[lang]["page" if len(nums) == 1 else "pages"]
            lines.append(f"{name} ({label} {', '.join(str(n) for n in sorted(nums))})")
    return lines
 
 
def plan_slides(
    deck: PresentationContent, include_sources_slide: bool = True, lang: str = "en"
) -> list[Planned]:
    """Decide the slides to draw: make sure there is a title slide, give the last text
    slide the 'closing' look, split long tables, and add a Sources slide at the end."""
    slides = sorted(deck.slides, key=lambda s: s.slide_number)
    if not slides or _layout_of(slides[0]) != "title":
        slides.insert(
            0,
            SlideContent(
                slide_number=0,
                title=deck.presentation_title,
                layout_type="title",
                subtitle=f"{LABELS[lang]['audience']} {deck.target_audience}",
            ),
        )
 
    planned: list[Planned] = []
    for index, s in enumerate(slides):
        layout = _layout_of(s)
        if layout == "bullet_points" and index == len(slides) - 1 and index > 0:
            layout = "closing"  # the last text slide is the conclusion
        if layout == "table" and len(s.table.rows) > MAX_TABLE_ROWS_PER_SLIDE:
            rows = s.table.rows
            for i in range(0, len(rows), MAX_TABLE_ROWS_PER_SLIDE):
                part = s.table.model_copy(
                    update={"rows": rows[i : i + MAX_TABLE_ROWS_PER_SLIDE]}
                )
                title = s.title if i == 0 else f"{s.title} {LABELS[lang]['cont']}"
                piece = s.model_copy(
                    update={
                        "table": part,
                        "title": title,
                        "bullets": s.bullets if i == 0 else [],
                        "speaker_notes": s.speaker_notes if i == 0 else "",
                    }
                )
                planned.append(Planned(piece, "table", _footer_text(piece, lang)))
        else:
            planned.append(Planned(s, layout, _footer_text(s, lang)))
 
    lines = _sources_lines(slides, lang)
    for slide in slides:
        for data in (slide.chart, slide.table):
            if data is not None and data.source:
                source = data.source.strip()
                if source and source not in lines:
                    lines.append(source)
    if include_sources_slide and lines:
        sources = SlideContent(
            slide_number=999,
            title=LABELS[lang]["sources"],
            layout_type="bullet_points",
            bullets=[{"text": line, "source_ids": []} for line in lines],
        )
        planned.append(Planned(sources, "sources"))
    return planned
 
 
# ---------------------------------------------------------------- slide frames
def _background(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(color)
 
 
def _content_slide(prs, p: Planned, th: Theme, rtl: bool, number: int):
    s = p.slide
    slide = prs.slides.add_slide(
        prs.slide_layouts[5]
    )  # "Title Only": a real title placeholder
    _background(slide, th.background)
    _rect(slide, MARGIN, 0.55, 0.1, 0.8, th.accent, rtl)
 
    title = slide.shapes.title
    w = SLIDE_W - 2 * MARGIN - 0.3
    title.left, title.top = Inches(_mx(MARGIN + 0.3, w, rtl)), Inches(0.4)
    title.width, title.height = Inches(w), Inches(1.1)
    tf = title.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    size = 34 if len(s.title) <= 55 else 28
    _para(tf, True, s.title, size, th.heading, th.font_heading, rtl, bold=True)
 
    # footer: slide number + sources
    _, nf = _textbox(slide, SLIDE_W - MARGIN - 0.8, 6.98, 0.8, 0.35, rtl)
    _para(
        nf,
        True,
        str(number),
        12,
        th.muted,
        th.font_body,
        rtl,
        align=PP_ALIGN.LEFT if rtl else PP_ALIGN.RIGHT,
    )
    footer = p.footer
    if footer:
        _, sf = _textbox(slide, MARGIN, 6.98, SLIDE_W - 2 * MARGIN - 1.0, 0.35, rtl)
        _para(sf, True, footer, 11, th.muted, th.font_body, rtl, italic=True)
 
    if s.speaker_notes:
        slide.notes_slide.notes_text_frame.text = s.speaker_notes
    return slide
 
 
# ---------------------------------------------------------------- layouts
def _render_title(prs, p, th, rtl, number):
    s = p.slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    _background(slide, th.title_background)
    _rect(slide, MARGIN, 3.55, 1.6, 0.08, th.accent2, rtl)
    w = SLIDE_W - 2 * MARGIN
    title, sub = slide.shapes.title, slide.placeholders[1]
    for shp, y, h in ((title, 1.6, 1.8), (sub, 3.85, 1.3)):
        shp.left, shp.top = Inches(_mx(MARGIN, w, rtl)), Inches(y)
        shp.width, shp.height = Inches(w), Inches(h)
        shp.text_frame.word_wrap = True
    title.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
    size = 46 if len(s.title) <= 50 else 36
    _para(
        title.text_frame,
        True,
        s.title,
        size,
        th.title_text,
        th.font_heading,
        rtl,
        bold=True,
    )
    _para(sub.text_frame, True, s.subtitle or "", 24, th.title_text, th.font_body, rtl)
    if s.speaker_notes:
        slide.notes_slide.notes_text_frame.text = s.speaker_notes
 
 
def _render_bullets(prs, p, th, rtl, number):
    s = p.slide
    slide = _content_slide(prs, p, th, rtl, number)
    texts = [b.text for b in s.bullets]
    _, tf = _textbox(slide, MARGIN + 0.1, BODY_TOP, SLIDE_W - 2 * MARGIN - 0.2, BODY_H, rtl)
    size = _bullet_size(texts)
    for i, text in enumerate(texts):
        _para(
            tf, i == 0, text, size, th.text, th.font_body, rtl, bullet=True, space_after=14
        )
 
 
def _render_closing(prs, p, th, rtl, number):
    s = p.slide
    slide = _content_slide(prs, p, th, rtl, number)
    texts = [b.text for b in s.bullets]
    _rect(slide, MARGIN, BODY_TOP, SLIDE_W - 2 * MARGIN, BODY_H, th.card, rtl)
    _rect(slide, MARGIN, BODY_TOP, 0.12, BODY_H, th.accent2, rtl)
    _, tf = _textbox(
        slide, MARGIN + 0.5, BODY_TOP + 0.3, SLIDE_W - 2 * MARGIN - 0.9, BODY_H - 0.6, rtl
    )
    size = _bullet_size(texts, base=26)
    for i, text in enumerate(texts):
        _para(
            tf, i == 0, text, size, th.text, th.font_body, rtl, bullet=True, space_after=16
        )
 
 
def _render_two_columns(prs, p, th, rtl, number):
    s = p.slide
    slide = _content_slide(prs, p, th, rtl, number)
    gap, col_w = 0.4, (SLIDE_W - 2 * MARGIN - 0.4) / 2
    columns = [[b.text for b in s.bullets], [b.text for b in s.right_bullets]]
    size = _bullet_size([t for col in columns for t in col], base=22)
    for i, texts in enumerate(columns):
        x = MARGIN + i * (col_w + gap)
        _rect(slide, x, BODY_TOP, col_w, BODY_H, th.card, rtl)
        _rect(slide, x, BODY_TOP, col_w, 0.1, th.accent if i == 0 else th.accent2, rtl)
        _, tf = _textbox(slide, x + 0.3, BODY_TOP + 0.35, col_w - 0.6, BODY_H - 0.6, rtl)
        for j, text in enumerate(texts):
            _para(
                tf,
                j == 0,
                text,
                size,
                th.text,
                th.font_body,
                rtl,
                bullet=True,
                space_after=12,
            )
 
 
def _render_chart(prs, p, th, rtl, number):
    s = p.slide
    slide = _content_slide(prs, p, th, rtl, number)
    texts = [b.text for b in s.bullets]
    full_w = SLIDE_W - 2 * MARGIN
    chart_w = full_w * 0.62 if texts else full_w
    pptx_charts.add_chart(
        slide,
        s.chart,
        th,
        Inches(_mx(MARGIN, chart_w, rtl)),
        Inches(BODY_TOP),
        Inches(chart_w),
        Inches(BODY_H),
    )
    if texts:
        x = MARGIN + chart_w + 0.35
        w = full_w - chart_w - 0.35
        _, tf = _textbox(slide, x, BODY_TOP + 0.2, w, BODY_H - 0.4, rtl)
        size = _bullet_size(texts, base=20)
        for i, text in enumerate(texts):
            _para(
                tf,
                i == 0,
                text,
                size,
                th.text,
                th.font_body,
                rtl,
                bullet=True,
                space_after=12,
            )
 
 
def _render_table(prs, p, th, rtl, number):
    s = p.slide
    slide = _content_slide(prs, p, th, rtl, number)
    t = s.table
    w = SLIDE_W - 2 * MARGIN
    top = BODY_TOP
    if t.title:  # a small caption above the table
        _, cf = _textbox(slide, MARGIN, top - 0.05, w, 0.4, rtl)
        _para(
            cf, True, t.title, 14, th.muted, th.font_body, rtl, italic=True, space_after=0
        )
        top += 0.45
    texts = [s_.text for s_ in s.bullets]
    bullets_h = 1.2 if texts else 0
    n_rows, n_cols = len(t.rows) + 1, len(t.columns)
    row_h = min(0.55, (BODY_TOP + BODY_H - top - bullets_h) / n_rows)
    frame = slide.shapes.add_table(
        n_rows,
        n_cols,
        Inches(_mx(MARGIN, w, rtl)),
        Inches(top),
        Inches(w),
        Inches(row_h * n_rows),
    )
    table = frame.table
    if rtl:
        frame._element.graphic.graphicData.tbl.tblPr.set("rtl", "1")
    size = 16 if n_cols <= 3 else 14 if n_cols <= 5 else 12
    for c in range(n_cols):
        table.columns[c].width = Inches(w / n_cols)
    for r in range(n_rows):
        table.rows[r].height = Inches(row_h)
        for c in range(n_cols):
            cell = table.cell(r, c)
            header = r == 0
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(
                th.primary if header else (th.card if r % 2 else th.background)
            )
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            value = t.columns[c] if header else t.rows[r - 1][c]
            numeric = not header and _is_number(value)
            align = (
                PP_ALIGN.LEFT if (numeric and rtl) else PP_ALIGN.RIGHT if numeric else None
            )
            _para(
                cell.text_frame,
                True,
                value if header else _fmt_cell(value),
                size,
                "FFFFFF" if header else th.text,
                th.font_body,
                rtl,
                bold=header,
                align=align,
                space_after=0,
            )
    if texts:
        y = top + row_h * n_rows + 0.2
        _, tf = _textbox(slide, MARGIN + 0.1, y, w - 0.2, BODY_TOP + BODY_H - y, rtl)
        for i, text in enumerate(texts):
            _para(
                tf, i == 0, text, 18, th.text, th.font_body, rtl, bullet=True, space_after=6
            )
 
 
_RENDERERS = {
    "title": _render_title,
    "bullet_points": _render_bullets,
    "sources": _render_bullets,
    "closing": _render_closing,
    "two_columns": _render_two_columns,
    "chart": _render_chart,
    "table": _render_table,
}
 
 
def build_pptx(
    deck: PresentationContent,
    out_path: str | Path,
    theme: str | Theme = "classic",
    rtl: bool | None = None,
    include_sources_slide: bool = True,
) -> Path:
    """Render `deck` to `out_path` (a .pptx file) and return the path.
 
    rtl=None means "decide from the text": Arabic or Hebrew decks are mirrored.
    """
    th = get_theme(theme) if isinstance(theme, str) else theme
    rtl = detect_rtl(deck) if rtl is None else rtl
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    prs.core_properties.title = deck.presentation_title
    lang = detect_language(deck)
    for number, planned in enumerate(
        plan_slides(deck, include_sources_slide, lang), start=1
    ):
        _RENDERERS[planned.kind](prs, planned, th, rtl, number)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    return out_path
