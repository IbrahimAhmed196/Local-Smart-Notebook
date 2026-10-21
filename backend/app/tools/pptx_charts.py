"""Native (editable) PowerPoint charts, drawn from the ChartData that Task 2 computed."""
 
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Pt
 
from app.schemas.slide_content import ChartData
from app.tools.pptx_themes import Theme
 
_KINDS = {
    "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
}
 
 
def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)
 
 
def _number_format(chart: ChartData) -> str:
    """Whole numbers get no decimals (200,000); anything else gets two (12.50)."""
    whole = all(float(v).is_integer() for s in chart.series for v in s.values)
    return "#,##0" if whole else "#,##0.00"
 
 
def add_chart(slide, chart: ChartData, theme: Theme, x, y, cx, cy):
    """Add an editable chart (in PowerPoint: right-click > Edit Data)."""
    data = CategoryChartData()
    data.categories = chart.categories
    for s in chart.series:
        data.add_series(s.name, s.values)
 
    frame = slide.shapes.add_chart(_KINDS[chart.chart_type], x, y, cx, cy, data)
    ch = frame.chart
    ch.font.size = Pt(14)
    ch.font.name = theme.font_body
    ch.font.color.rgb = _rgb(theme.text)
 
    if chart.title:
        ch.has_title = True
        ch.chart_title.text_frame.text = chart.title
        run = ch.chart_title.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = _rgb(theme.heading)
    else:
        ch.has_title = False
 
    many_series = len(chart.series) > 1
    ch.has_legend = many_series or chart.chart_type == "pie"
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
 
    plot = ch.plots[0]
    fmt = _number_format(chart)
    if chart.chart_type == "pie":
        plot.has_data_labels = True
        labels = plot.data_labels
        labels.font.size = Pt(14)
        labels.font.color.rgb = _rgb("FFFFFF")
        labels.show_percentage = True
        labels.show_value = False
        for i, point in enumerate(plot.series[0].points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = _rgb(theme.palette[i % len(theme.palette)])
        return frame
 
    for i, series in enumerate(plot.series):
        color = _rgb(theme.palette[i % len(theme.palette)])
        if chart.chart_type == "bar":
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = color
        else:
            series.format.line.color.rgb = color
            series.format.line.width = Pt(3)
            series.smooth = False
    if not many_series and chart.chart_type == "bar":
        plot.has_data_labels = True
        plot.data_labels.font.size = Pt(12)
        plot.data_labels.number_format = fmt
        plot.data_labels.number_format_is_linked = False
    ch.value_axis.tick_labels.number_format = fmt
    ch.value_axis.tick_labels.number_format_is_linked = False
    ch.value_axis.has_major_gridlines = True
    ch.value_axis.major_gridlines.format.line.color.rgb = _rgb(theme.muted)
    ch.value_axis.major_gridlines.format.line.width = Pt(0.25)
    return frame
